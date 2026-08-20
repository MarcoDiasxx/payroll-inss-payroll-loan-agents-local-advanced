from __future__ import annotations

import truststore

truststore.inject_into_ssl()

import json
import os
from pathlib import Path
from typing import Any

from dotenv import load_dotenv
from fastmcp import FastMCP
from openpyxl import load_workbook

from esteiras_mcp.atlassian import AtlassianClient

load_dotenv()
mcp = FastMCP("Esteiras INSS - Jira e Confluence")
ROOT = Path(__file__).resolve().parents[2]


def split_env(name: str) -> list[str]:
    return [x.strip() for x in os.getenv(name, "").split(",") if x.strip()]


def client() -> AtlassianClient:
    return AtlassianClient()


def taxonomy() -> dict[str, Any]:
    return json.loads((ROOT / "config/domain-taxonomy.json").read_text(encoding="utf-8"))


@mcp.tool
def test_atlassian_access() -> dict[str, Any]:
    """Testa a identidade Atlassian configurada no .env, sem expor segredos."""
    return {"ok": True, "identity": client().myself(), "readOnly": True}


@mcp.tool
def get_jira_issue(issue_key: str) -> dict[str, Any]:
    """Abre um card Jira por chave, incluindo descricao, comentarios e anexos listados."""
    return client().issue(issue_key)


@mcp.tool
def search_jira_esteiras(query: str, max_results: int = 30) -> list[dict[str, Any]]:
    """Pesquisa cards Jira relacionados a esteiras. Use termos de negocio, siglas e sinonimos."""
    limit = min(max_results, int(os.getenv("SEARCH_MAX_RESULTS", "50")))
    return client().search_jira(query, limit, split_env("JIRA_PROJECT_KEYS"))


@mcp.tool
def get_confluence_page(page_id: str) -> dict[str, Any]:
    """Abre uma pagina do Confluence pelo ID e retorna texto, versao e URL."""
    return client().confluence_page(page_id)


@mcp.tool
def search_confluence_esteiras(query: str, max_results: int = 30) -> list[dict[str, Any]]:
    """Pesquisa paginas no Confluence relacionadas a esteiras e regras de negocio."""
    limit = min(max_results, int(os.getenv("SEARCH_MAX_RESULTS", "50")))
    return client().search_confluence(query, limit, split_env("CONFLUENCE_SPACE_KEYS"))


@mcp.tool
def discover_domain_knowledge(domain: str) -> dict[str, Any]:
    """Faz busca ampla Jira+Confluence para INSS, margem nova, refin, portabilidade e derivados."""
    tx = taxonomy()
    terms = tx["domains"].get(domain, [domain])
    expanded = list(dict.fromkeys(terms + tx["cross_cutting"]))
    # Keep calls bounded while covering business aliases.
    queries = expanded[:10]
    jira: list[dict[str, Any]] = []
    conf: list[dict[str, Any]] = []
    seen_jira: set[str] = set()
    seen_conf: set[str] = set()
    c = client()
    for q in queries:
        for item in c.search_jira(q, 15, split_env("JIRA_PROJECT_KEYS")):
            if item["key"] not in seen_jira:
                seen_jira.add(item["key"]); jira.append(item)
        for item in c.search_confluence(q, 15, split_env("CONFLUENCE_SPACE_KEYS")):
            if item["id"] not in seen_conf:
                seen_conf.add(item["id"]); conf.append(item)
    return {"domain": domain, "queries": queries, "jira": jira, "confluence": conf}


@mcp.tool
def get_core_references() -> dict[str, Any]:
    """Carrega todas as referencias core absolutas e informa qualquer fonte indisponivel."""
    c = client()
    errors: list[dict[str, str]] = []
    jira = []
    confluence = []
    for key in split_env("JIRA_CORE_ISSUES"):
        try: jira.append(c.issue(key))
        except Exception as exc: errors.append({"source": f"jira:{key}", "error": str(exc)})
    for page_id in split_env("CONFLUENCE_CORE_PAGE_IDS"):
        try: confluence.append(c.confluence_page(page_id))
        except Exception as exc: errors.append({"source": f"confluence:{page_id}", "error": str(exc)})
    policy_path = ROOT / "docs/core/esteiras-core-policy.md"
    workbook = Path(os.getenv("OFFICIAL_WORKBOOK_LOCAL_PATH", ""))
    if not workbook.is_absolute(): workbook = ROOT / workbook
    return {
        "absolute": True,
        "jira": jira,
        "confluence": confluence,
        "policy": {"path": str(policy_path), "available": policy_path.is_file(), "content": policy_path.read_text(encoding="utf-8") if policy_path.is_file() else None},
        "workbook": {"path": str(workbook), "available": workbook.is_file(), "sharepointUrl": os.getenv("SHAREPOINT_CORE_WORKBOOK_URL")},
        "errors": errors,
        "complete": bool(jira and confluence and policy_path.is_file() and workbook.is_file() and not errors)
    }


@mcp.tool
def lookup_official_activity(decision_number: int | None = None, text: str | None = None) -> dict[str, Any]:
    """Consulta a planilha oficial local por numero de decisao ou texto de atividade."""
    raw = os.getenv("OFFICIAL_WORKBOOK_LOCAL_PATH", "")
    path = Path(raw)
    if not path.is_absolute(): path = ROOT / path
    if not path.is_file():
        return {"ok": False, "blocked": True, "reason": "Planilha oficial local indisponivel", "path": str(path)}
    wb = load_workbook(path, read_only=True, data_only=True)
    matches = []
    needle = (text or "").casefold()
    for ws in wb.worksheets:
        for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            values = ["" if value is None else str(value) for value in row]
            hit_num = decision_number is not None and any(v.strip() == str(decision_number) for v in values)
            hit_text = bool(needle) and any(needle in v.casefold() for v in values)
            if hit_num or hit_text:
                matches.append({"sheet": ws.title, "row": row_idx, "values": values})
                if len(matches) >= 100: break
        if len(matches) >= 100: break
    return {"ok": True, "path": str(path), "matches": matches}


if __name__ == "__main__":
    mcp.run()
